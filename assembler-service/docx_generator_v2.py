from __future__ import annotations

import argparse
import io
import json
import re
import sys
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# GigaChat часто отдаёт Markdown — в DOCX/PDF это должно стать настоящим форматированием.
_MD_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_MD_ITALIC_RE = re.compile(r"(?<!\*)\*([^*\n]+)\*(?!\*)")
_MD_HEADING_RE = re.compile(r"^\s{0,3}(#{1,6})\s+(.*)$")
_MD_HR_RE = re.compile(r"^\s{0,3}([-*_])\1{2,}\s*$")


class DocumentValidationError(ValueError):
    """Входной JSON не соответствует ожидаемой схеме."""


# Ожидаемый формат JSON
# {
#   "document_id": "b6b6b6b6-...",
#   "station_name": "Чита-1",
#   "region": "chita",
#   "organization": "АО «Чита Главснаб»",
#   "generated_at": "2026-07-13T10:00:00Z",
#   "sections": [
#       {
#           "id": "0001", "order": 1, "title": "Общие сведения", "text": "...",
#           "tables": [
#               {"title": "Показатели", "headers": ["Параметр", "Значение"],
#                "rows": [["Длина пути", "850 м"], ["Кол-во путей", "12"]]}
#           ]
#       },
#       {"id": "0002", "order": 2, "title": "Технические характеристики", "text": "..."}
#   ],
#   "appendices": [
#       {"title": "Схема путей станции", "type": "image", "source_url": "s3://.../scheme.png"},
#       {"title": "Перечень примыканий", "type": "text", "text": "..."},
#       {"title": "Ведомость путей", "type": "table",
#        "table": {"headers": ["Путь", "Длина, м", "Назначение"],
#                  "rows": [["1", "850", "Приёмо-отправочный"], ["2", "1050", "Сортировочный"]]}}
#   ]
# }

@dataclass
class Table:
    headers: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    title: Optional[str] = None

    @classmethod
    def from_dict(cls, data: dict[str, Any], context: str) -> "Table":
        rows_raw = data.get("rows")
        if not isinstance(rows_raw, list) or not rows_raw:
            raise DocumentValidationError(f"{context}: поле 'rows' должно быть непустым списком")
        headers_raw = data.get("headers", [])
        if headers_raw and not isinstance(headers_raw, list):
            raise DocumentValidationError(f"{context}: поле 'headers' должно быть списком")

        n_cols = len(headers_raw) if headers_raw else len(rows_raw[0])
        rows: list[list[str]] = []
        for r_idx, row in enumerate(rows_raw):
            if not isinstance(row, list):
                raise DocumentValidationError(f"{context}: строка #{r_idx} таблицы должна быть списком")
            cells = [str(c) for c in row]
            if len(cells) < n_cols:
                cells += [""] * (n_cols - len(cells))
            elif len(cells) > n_cols:
                raise DocumentValidationError(
                    f"{context}: строка #{r_idx} содержит {len(cells)} колонок, ожидалось {n_cols}"
                )
            rows.append(cells)

        return cls(
            headers=[str(h) for h in headers_raw],
            rows=rows,
            title=data.get("title"),
        )


@dataclass
class Section:
    id: str
    order: int
    title: str
    text: str
    tables: list[Table] = field(default_factory=list)


@dataclass
class Appendix:
    title: str
    type: str  # "image" | "text" | "table"
    text: Optional[str] = None
    source_url: Optional[str] = None
    table: Optional[Table] = None


@dataclass
class InstructionDocument:
    document_id: str
    station_name: str
    region: str
    organization: Optional[str]
    generated_at: str
    sections: list[Section]
    appendices: list[Appendix] = field(default_factory=list)

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> "InstructionDocument":
        required_top_level = ("document_id", "station_name", "region", "sections")
        missing = [key for key in required_top_level if key not in data]
        if missing:
            raise DocumentValidationError(
                f"В JSON отсутствуют обязательные поля: {', '.join(missing)}"
            )

        raw_sections = data["sections"]
        if not isinstance(raw_sections, list) or not raw_sections:
            raise DocumentValidationError("Поле 'sections' должно быть непустым списком")

        sections: list[Section] = []
        for i, raw in enumerate(raw_sections):
            for key in ("title", "text"):
                if key not in raw:
                    raise DocumentValidationError(
                        f"Секция #{i}: отсутствует обязательное поле '{key}'"
                    )
            raw_tables = raw.get("tables", [])
            if not isinstance(raw_tables, list):
                raise DocumentValidationError(f"Секция #{i}: поле 'tables' должно быть списком")
            tables = [
                Table.from_dict(t, context=f"Секция #{i}, таблица #{t_idx}")
                for t_idx, t in enumerate(raw_tables)
            ]
            sections.append(
                Section(
                    id=str(raw.get("id", i)),
                    order=int(raw.get("order", i)),
                    title=str(raw["title"]).strip(),
                    text=str(raw["text"]).strip(),
                    tables=tables,
                )
            )
        sections.sort(key=lambda s: s.order)

        appendices: list[Appendix] = []
        for i, a in enumerate(data.get("appendices", [])):
            if "title" not in a:
                raise DocumentValidationError(f"Приложение #{i}: отсутствует поле 'title'")
            a_type = str(a.get("type", "text"))
            table_obj = None
            if a_type == "table":
                if "table" not in a:
                    raise DocumentValidationError(
                        f"Приложение #{i} ('{a['title']}'): для type='table' требуется поле 'table'"
                    )
                table_obj = Table.from_dict(a["table"], context=f"Приложение #{i}")
            appendices.append(
                Appendix(
                    title=str(a["title"]),
                    type=a_type,
                    text=a.get("text"),
                    source_url=a.get("source_url"),
                    table=table_obj,
                )
            )

        return cls(
            document_id=str(data["document_id"]),
            station_name=str(data["station_name"]),
            region=str(data["region"]),
            organization=data.get("organization"),
            generated_at=data.get("generated_at", datetime.now(timezone.utc).isoformat()),
            sections=sections,
            appendices=appendices,
        )


# Шаблон описывает оформление.
# Помимо базовых ключей поддерживаются (со значениями по умолчанию в _TEMPLATE_DEFAULTS):
#   margins_cm: (top, bottom, left, right) в сантиметрах
#   line_spacing_pt: точный межстрочный интервал в пунктах (None -> множитель 1.15)
#   first_line_indent_cm: абзацный отступ первой строки тела текста
#   heading_prefix: если задан (например "РАЗДЕЛ"), заголовки разделов формируются как
#       "{prefix} {N}.   {TITLE В ВЕРХНЕМ РЕГИСТРЕ}", по центру, жирным, без цвета;
#       если None — прежний формат "{N}. {Title}" слева, жирным, цветом акцента.
#   header_enabled: показывать ли текст в верхнем колонтитуле (название станции)
#   table_header_shading: цвет заливки шапки таблицы (RRGGBB); "__accent__" — взять
#       accent_color; None — без заливки (только жирный текст и границы)
REGION_TEMPLATES: dict[str, dict[str, Any]] = {
    "default": {
        "font_name": "Times New Roman",
        "base_font_size": 12,
        "heading_font_size": 14,
        "title_font_size": 20,
        "accent_color": "000000",  # RRGGBB
        "title_prefix": "Инструкция по эксплуатации станции",
        "footer_text": "Документ сформирован автоматически",
    },
    "chita": {
        # Оформление приведено в соответствие с типовым ГОСТ-документом
        # «Инструкция о порядке обслуживания и организации движения на
        # железнодорожном пути необщего пользования» Забайкальской ж.д.:
        # Times New Roman 14pt, чёрный текст без цветных акцентов, разделы
        # оформлены как "РАЗДЕЛ N.  НАЗВАНИЕ" по центру, абзацный отступ
        # первой строки 1,25 см, точный интервал 18пт, поля 2/2/2,5/1,5 см,
        # колонтитул — только номер страницы, без надписей.
        "font_name": "Times New Roman",
        "base_font_size": 14,
        "heading_font_size": 14,
        "title_font_size": 16,
        "accent_color": "000000",
        "title_prefix": [
            "ИНСТРУКЦИЯ",
            "о порядке обслуживания и организации движения",
            "на железнодорожном пути необщего пользования",
        ],
        "footer_text": "",
        "margins_cm": (2.0, 2.0, 2.5, 1.5),
        "line_spacing_pt": 18,
        "first_line_indent_cm": 1.25,
        "heading_prefix": "РАЗДЕЛ",
        "header_enabled": False,
        "table_header_shading": None,
    },
}

_TEMPLATE_DEFAULTS: dict[str, Any] = {
    "margins_cm": (56 / 28.3465, 56 / 28.3465, 85 / 28.3465, 56 / 28.3465),
    "line_spacing_pt": None,
    "first_line_indent_cm": 0.0,
    "heading_prefix": None,
    "header_enabled": True,
    "table_header_shading": "__accent__",
}


def get_template(region: str) -> dict[str, Any]:
    base = dict(_TEMPLATE_DEFAULTS)
    base.update(REGION_TEMPLATES.get(region, REGION_TEMPLATES["default"]))
    return base


# Сборка DOCX
class InstructionDocxBuilder:
    """Собирает объект docx.Document из InstructionDocument по шаблону региона."""

    def __init__(self, doc_data: InstructionDocument):
        self.data = doc_data
        self.template = get_template(doc_data.region)
        self.document = Document()

    def build(self) -> Document:
        self._setup_base_style()
        self._add_header_footer()
        self._add_title_page()
        self._add_sections()
        if self.data.appendices:
            self._add_appendices()
        return self.document

    # Базовый стиль
    def _setup_base_style(self) -> None:
        style = self.document.styles["Normal"]
        style.font.name = self.template["font_name"]
        style.font.size = Pt(self.template["base_font_size"])
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), self.template["font_name"])

        top_cm, bottom_cm, left_cm, right_cm = self.template["margins_cm"]
        section = self.document.sections[0]
        section.top_margin = Cm(top_cm)
        section.bottom_margin = Cm(bottom_cm)
        section.left_margin = Cm(left_cm)
        section.right_margin = Cm(right_cm)

    # Колонтитулы
    def _add_header_footer(self) -> None:
        section = self.document.sections[0]

        if self.template["header_enabled"]:
            header_p = section.header.paragraphs[0]
            header_p.text = f"Станция {self.data.station_name}"
            header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for run in header_p.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor.from_string("808080")

        footer_text = self.template["footer_text"]
        footer_p = section.footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if not footer_text else WD_ALIGN_PARAGRAPH.CENTER
        footer_p.text = f"{footer_text}  ·  стр. " if footer_text else ""
        for run in footer_p.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string("808080")
        self._add_page_number_field(footer_p)

    @staticmethod
    def _add_page_number_field(paragraph) -> None:
        """Вставляет динамическое поле номера страницы (PAGE) в параграф."""
        run = paragraph.add_run()
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = "PAGE"
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_end)

    # Титульный блок
    def _add_title_page(self) -> None:
        if isinstance(self.template["title_prefix"], list):
            self._add_title_page_gost()
        else:
            self._add_title_page_default()

    def _add_title_page_default(self) -> None:
        accent = RGBColor.from_string(self.template["accent_color"])

        p = self.document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(self.template["title_prefix"])
        run.bold = True
        run.font.size = Pt(self.template["title_font_size"])
        run.font.color.rgb = accent

        p2 = self.document.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run2 = p2.add_run(self.data.station_name)
        run2.bold = True
        run2.font.size = Pt(self.template["title_font_size"] + 4)
        run2.font.color.rgb = accent

        if self.data.organization:
            p3 = self.document.add_paragraph()
            p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run3 = p3.add_run(self.data.organization)
            run3.italic = True
            run3.font.size = Pt(12)

        meta = self.document.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        gen_date = self._format_date(self.data.generated_at)
        meta_run = meta.add_run(f"Документ №{self.data.document_id}  ·  {gen_date}")
        meta_run.font.size = Pt(9)
        meta_run.font.color.rgb = RGBColor.from_string("808080")

        self.document.add_page_break()

    def _add_title_page_gost(self) -> None:
        """ГОСТ-титул на одной странице: заголовок сверху/по центру, год — в самом низу."""
        title_lines = list(self.template["title_prefix"])
        subject_line = (
            f"{self.data.organization}, примыкающем к железнодорожной станции {self.data.station_name}"
            if self.data.organization
            else f"примыкающем к железнодорожной станции {self.data.station_name}"
        )

        section = self.document.sections[0]
        # Запас под колонтитул + служебный абзац после таблицы (иначе таблица
        # выталкивает пустую 2-ю страницу перед контентом).
        usable_height = (
            section.page_height - section.top_margin - section.bottom_margin - Cm(2.2)
        )
        year_height = Cm(1.6)
        title_height = usable_height - year_height
        if int(title_height) < int(Cm(14)):
            title_height = Cm(14)

        table = self.document.add_table(rows=2, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        top_row, bottom_row = table.rows[0], table.rows[1]
        top_row.height = title_height
        top_row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        bottom_row.height = year_height
        bottom_row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY

        top_cell, bottom_cell = top_row.cells[0], bottom_row.cells[0]
        top_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        bottom_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        self._clear_table_borders(table)

        # Заголовок в верхней (высокой) ячейке
        top_cell.text = ""
        first_p = top_cell.paragraphs[0]
        first_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run0 = first_p.add_run(title_lines[0] if title_lines else "ИНСТРУКЦИЯ")
        run0.bold = True
        run0.font.size = Pt(self.template["title_font_size"])

        for line in title_lines[1:]:
            p = top_cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line)
            run.bold = True
            run.font.size = Pt(self.template["title_font_size"])

        subj_p = top_cell.add_paragraph()
        subj_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subj_p.paragraph_format.space_before = Pt(12)
        subj_run = subj_p.add_run(subject_line)
        subj_run.bold = True
        subj_run.font.size = Pt(self.template["title_font_size"])

        # Год — нижняя ячейка = низ первой страницы
        gen_date = self._format_date(self.data.generated_at)
        year = gen_date.split(".")[-1] if "." in gen_date else gen_date
        bottom_cell.text = ""
        year_p = bottom_cell.paragraphs[0]
        year_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        year_run = year_p.add_run(f"{year} г")
        year_run.bold = True
        year_run.font.size = Pt(self.template["base_font_size"])

        # Разрыв страницы в служебном абзаце СРАЗУ после таблицы —
        # без отдельного add_page_break() (он давал пустую страницу).
        self._page_break_after_table(table)

    @staticmethod
    def _clear_table_borders(table) -> None:
        tbl = table._tbl
        tbl_pr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
        if tbl.tblPr is None:
            tbl.insert(0, tbl_pr)
        borders = tbl_pr.find(qn("w:tblBorders"))
        if borders is not None:
            tbl_pr.remove(borders)
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            elem = OxmlElement(f"w:{edge}")
            elem.set(qn("w:val"), "nil")
            elem.set(qn("w:sz"), "0")
            elem.set(qn("w:space"), "0")
            elem.set(qn("w:color"), "auto")
            borders.append(elem)
        tbl_pr.append(borders)

    def _page_break_after_table(self, table) -> None:
        """Ставит page-break в абзац сразу после таблицы (он всегда есть в python-docx)."""
        from docx.text.paragraph import Paragraph

        tbl = table._tbl
        next_el = tbl.getnext()
        if next_el is not None and next_el.tag == qn("w:p"):
            paragraph = Paragraph(next_el, table._parent)
            # очистить служебный абзац и оставить только разрыв страницы
            p_element = paragraph._p
            for child in list(p_element):
                if child.tag != qn("w:pPr"):
                    p_element.remove(child)
            run = paragraph.add_run()
            br = OxmlElement("w:br")
            br.set(qn("w:type"), "page")
            run._r.append(br)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            return
        self.document.add_page_break()

    @staticmethod
    def _format_date(iso_string: str) -> str:
        try:
            dt = datetime.fromisoformat(iso_string.replace("Z", "+00:00"))
            return dt.strftime("%d.%m.%Y")
        except ValueError:
            return iso_string

    # Разделы
    def _add_sections(self) -> None:
        accent = RGBColor.from_string(self.template["accent_color"])
        heading_prefix = self.template["heading_prefix"]

        for i, section in enumerate(self.data.sections, start=1):
            heading = self.document.add_paragraph()
            if heading_prefix:
                heading_text = f"{heading_prefix} {i}.   {section.title.upper()}"
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                heading_text = f"{i}. {section.title}"
            heading_run = heading.add_run(heading_text)
            heading_run.bold = True
            heading_run.font.size = Pt(self.template["heading_font_size"])
            if not heading_prefix:
                heading_run.font.color.rgb = accent
            heading.paragraph_format.space_before = Pt(18)
            heading.paragraph_format.space_after = Pt(6)
            heading.paragraph_format.keep_with_next = True
            self._apply_line_spacing(heading.paragraph_format)

            self._add_section_body(section.text)

            for table in section.tables:
                self._add_table(table)

    def _add_section_body(self, text: str) -> None:
        """Пишет тело раздела: Markdown (`**`, `#`/`##`/`###`) → реальное оформление DOCX."""
        if not text or not text.strip():
            return

        # Нормализуем: иногда LLM отдаёт литералы \\n или смешивает \\r
        normalized = (
            text.replace("\\r\\n", "\n")
            .replace("\\n", "\n")
            .replace("\r\n", "\n")
            .replace("\r", "\n")
        )
        # Если почти нет переносов, но есть markdown-заголовки — режем перед ними
        if normalized.count("\n") < 2 and ("##" in normalized or "**" in normalized):
            normalized = re.sub(r"\s*(#{1,6}\s+)", r"\n\1", normalized)
            normalized = re.sub(r"\s+(\*\*\d)", r"\n\1", normalized)

        for raw_line in normalized.split("\n"):
            line = raw_line.strip()
            if not line or _MD_HR_RE.match(line):
                continue

            # Убрать ведущие markdown-заголовки даже если после # нет «чистого» match
            heading = _MD_HEADING_RE.match(line)
            if heading:
                self._add_inline_heading(heading.group(2).strip(), level=len(heading.group(1)))
                continue
            if line.startswith("#"):
                line = re.sub(r"^#{1,6}\s*", "", line).strip()
                if line:
                    self._add_inline_heading(line, level=3)
                continue

            only_bold = _MD_BOLD_RE.fullmatch(line)
            if only_bold:
                self._add_inline_heading(only_bold.group(1).strip(), level=3)
                continue

            self._add_body_paragraph(line)

    def _add_inline_heading(self, title: str, level: int = 2) -> None:
        title = _MD_BOLD_RE.sub(r"\1", title).strip()
        if not title:
            return
        p = self.document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10 if level <= 2 else 8)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.keep_with_next = True
        self._apply_line_spacing(p.paragraph_format)
        run = p.add_run(title)
        run.bold = True
        size = self.template["heading_font_size"] if level <= 2 else self.template["base_font_size"]
        run.font.size = Pt(size)

    def _add_body_paragraph(self, text: str) -> None:
        cleaned = "\n".join(
            _MD_HEADING_RE.sub(lambda m: m.group(2), line) if _MD_HEADING_RE.match(line) else line
            for line in text.split("\n")
        ).strip()
        # Страховка: одиночные # в начале строки
        cleaned = re.sub(r"^#{1,6}\s+", "", cleaned).strip()
        if not cleaned:
            return

        p = self.document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(6)
        indent_cm = self.template["first_line_indent_cm"]
        if indent_cm:
            p.paragraph_format.first_line_indent = Cm(indent_cm)
        self._apply_line_spacing(p.paragraph_format)
        self._append_markdown_runs(p, cleaned)
        # Если после парсера в абзаце всё ещё остались ** — вычистить runs
        for run in p.runs:
            if run.text and ("**" in run.text or run.text.startswith("#")):
                run.text = run.text.replace("**", "").lstrip("#").strip()

    @staticmethod
    def _append_markdown_runs(paragraph, text: str) -> None:
        """Разбивает текст на runs: **bold** и *italic*, маркеры в документ не попадают."""
        # Сначала bold, внутри/рядом — italic на оставшихся кусках
        pos = 0
        for match in _MD_BOLD_RE.finditer(text):
            if match.start() > pos:
                InstructionDocxBuilder._append_italic_runs(paragraph, text[pos:match.start()])
            run = paragraph.add_run(match.group(1))
            run.bold = True
            pos = match.end()
        if pos < len(text):
            InstructionDocxBuilder._append_italic_runs(paragraph, text[pos:])

    @staticmethod
    def _append_italic_runs(paragraph, text: str) -> None:
        if not text:
            return
        pos = 0
        for match in _MD_ITALIC_RE.finditer(text):
            if match.start() > pos:
                paragraph.add_run(text[pos:match.start()])
            run = paragraph.add_run(match.group(1))
            run.italic = True
            pos = match.end()
        if pos < len(text):
            paragraph.add_run(text[pos:])

    def _apply_line_spacing(self, paragraph_format) -> None:
        """Точный интервал (в пунктах) из шаблона, либо множитель 1.15 по умолчанию."""
        spacing_pt = self.template["line_spacing_pt"]
        if spacing_pt:
            paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            paragraph_format.line_spacing = Pt(spacing_pt)
        else:
            paragraph_format.line_spacing = 1.15

    # Приложения
    def _add_appendices(self) -> None:
        self.document.add_page_break()
        heading = self.document.add_paragraph()
        if self.template["heading_prefix"]:
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = heading.add_run("Приложения")
        run.bold = True
        run.font.size = Pt(self.template["heading_font_size"] + 2)
        if not self.template["heading_prefix"]:
            run.font.color.rgb = RGBColor.from_string(self.template["accent_color"])

        for i, appendix in enumerate(self.data.appendices, start=1):
            title_p = self.document.add_paragraph()
            title_run = title_p.add_run(f"Приложение {i}. {appendix.title}")
            title_run.bold = True
            title_p.paragraph_format.space_before = Pt(12)

            if appendix.type == "text" and appendix.text:
                self.document.add_paragraph(appendix.text)
            elif appendix.type == "image" and appendix.source_url:
                note = self.document.add_paragraph()
                note_run = note.add_run(f"[изображение: {appendix.source_url}]")
                note_run.italic = True
                note_run.font.color.rgb = RGBColor.from_string("808080")
            elif appendix.type == "table" and appendix.table:
                self._add_table(appendix.table)
            else:
                self.document.add_paragraph("(содержимое приложения недоступно)")

    # Таблицы
    def _add_table(self, table_data: Table) -> None:
        shading = self.template["table_header_shading"]
        shading = self.template["accent_color"] if shading == "__accent__" else shading

        if table_data.title:
            caption = self.document.add_paragraph()
            caption_run = caption.add_run(table_data.title)
            caption_run.italic = True
            caption_run.font.size = Pt(self.template["base_font_size"] - 1)
            caption.paragraph_format.space_before = Pt(6)
            caption.paragraph_format.space_after = Pt(4)

        n_cols = len(table_data.headers) if table_data.headers else len(table_data.rows[0])
        has_header = bool(table_data.headers)
        n_rows = len(table_data.rows) + (1 if has_header else 0)

        table = self.document.add_table(rows=n_rows, cols=n_cols)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        row_offset = 0
        if has_header:
            for c, header_text in enumerate(table_data.headers):
                cell = table.rows[0].cells[c]
                cell.text = ""
                p = cell.paragraphs[0]
                run = p.add_run(header_text)
                run.bold = True
                run.font.size = Pt(self.template["base_font_size"] - 1)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if shading:
                    run.font.color.rgb = RGBColor.from_string("FFFFFF")
                    self._set_cell_shading(cell, shading)
            row_offset = 1

        for r, row_values in enumerate(table_data.rows):
            for c, value in enumerate(row_values):
                cell = table.rows[r + row_offset].cells[c]
                cell.text = ""
                p = cell.paragraphs[0]
                run = p.add_run(value)
                run.font.size = Pt(self.template["base_font_size"] - 1)

        spacer = self.document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(6)

    @staticmethod
    def _set_cell_shading(cell, color_hex: str) -> None:
        """Заливка ячейки цветом (w:shd). ShadingType.SOLID даёт чёрный фон — используем CLEAR."""
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), color_hex)
        tc_pr.append(shd)


def generate_docx_bytes(payload: dict[str, Any] | str) -> bytes:
    """
    Основная точка входа модуля.
    """
    if isinstance(payload, str):
        payload = json.loads(payload)

    doc_data = InstructionDocument.from_dict(payload)
    document = InstructionDocxBuilder(doc_data).build()

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def generate_docx_file(payload: dict[str, Any] | str, output_path: str) -> str:
    """Обёртка над generate_docx_bytes для сохранения сразу на диск."""
    data = generate_docx_bytes(payload)
    with open(output_path, "wb") as f:
        f.write(data)
    return output_path


def _default_output_path(input_path: Path) -> Path:
    return input_path.with_suffix(".docx")


def main(argv: Optional[list[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        description="Генерация .docx инструкции по эксплуатации станции из JSON-файла."
    )
    parser.add_argument(
        "input_json",
        type=Path,
        help="Путь к JSON-файлу с данными документа (сегменты, приложения, таблицы).",
    )
    parser.add_argument(
        "-o", "--output",
        type=Path,
        default=None,
        help="Путь к выходному .docx (по умолчанию — рядом со входным файлом, с расширением .docx).",
    )
    args = parser.parse_args(argv)

    if not args.input_json.exists():
        parser.error(f"Файл не найден: {args.input_json}")

    try:
        payload = json.loads(args.input_json.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        print(f"Ошибка разбора JSON в {args.input_json}: {e}", file=sys.stderr)
        return 1

    output_path = args.output or _default_output_path(args.input_json)

    try:
        generate_docx_file(payload, str(output_path))
    except DocumentValidationError as e:
        print(f"Ошибка валидации данных: {e}", file=sys.stderr)
        return 1

    print(f"Документ сохранён: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())